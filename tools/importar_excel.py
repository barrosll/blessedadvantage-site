# Importa a calculadora de preços (Excel) para o catálogo da loja Welabb.
# Regras combinadas com o utilizador:
#  - preço = coluna K "Preço Final" arredondado para o ,90 mais próximo, mínimo 2,90 €
#  - os 8 produtos já no site recebem nome + preço (mantêm fotos e visibilidade)
#  - os 47 novos são criados; produtos com marcas registadas ou sem foto ficam escondidos
#  - fotos: capa do designer guardada no .3mf (D:\videos\STL_venda); MakerWorld bloqueia downloads
import io, json, math, os, re, sys, zipfile
from openpyxl import load_workbook
from PIL import Image

XLS = r"D:/calculadora_precos_impressao_3d-original(Recuperado Automaticamente).xlsx"
STL = r"D:/videos/STL_venda"
SITE = r"C:/Users/LOHAN/blessedadvantage-site"
base_catalog, out_catalog, out_report = sys.argv[1], sys.argv[2], sys.argv[3]

def preco(k):
    """Arredonda para o X,90 mais próximo; mínimo 2,90 €. Devolve cêntimos."""
    lo = math.floor(k - 0.9) + 0.9
    hi = lo + 1
    p = lo if (k - lo) < (hi - k) else hi
    return int(round(max(p, 2.9) * 100))

# linha da folha -> produto já existente no site
EXISTENTES = {7: ("produto-04", "Logo F1"), 25: ("produto-05", "Suporte VR caveira para óculos Meta"),
              27: ("produto-06", "Porta-óculos raposa"), 57: ("produto-07", "Anjo da guarda"),
              52: ("produto-08", "Batman FlowWeave"), 51: ("produto-09", "Fénix Voronoi"),
              19: ("produto-10", "Suporte de telemóvel"), 12: ("produto-11", "Pega para sacos")}

# linha -> (nome PT-PT, categoria, marca registada?, ficheiro .3mf com a foto ou None)
NOVOS = {
    4: ("Stitch 23 cm", "decoracao", True, "Re-plate Stitch A1.3mf"),
    5: ("Porta-cromos robusto Mundial 2026", "casa", True, "Porta_Figuritas_Copa_del_Mundo_2026(2).3mf"),
    6: ("Porta-chaves Seleções", "porta-chaves", True, "Llavero_Copa_del_Mundo_2026_Fifa_World_Cup.3mf"),
    8: ("Calendário F1 robusto", "decoracao", True, "F1_Track_Painted_V2.3mf"),
    9: ("Suporte de auscultadores F1", "casa", True, "f1_riginal(2).3mf"),
    10: ("Mealheiro Brasil", "presentes", True, "Brasilien_Spardose3mf.3mf"),
    11: ("Calendário MotoGP 2026 com suporte", "decoracao", True, "2026_Moto_GP_calendar.3mf"),
    13: ("Suporte Neymar Brasil", "decoracao", True, "brazil_neymar_jr_display(2).3mf"),
    14: ("Batman gordo", "decoracao", True, "Colored_BIG_Man.3mf"),
    15: ("Quadro Deadpool", "decoracao", True, "Deadpool_Graffiti_Front_200x200.3mf"),
    16: ("Caixa para câmara DJI Action 5 Pro", "casa", False, "DJI_ACTION_5_PRO_+_frame(2).3mf"),
    17: ("Réplica da taça do Mundial", "decoracao", True, "FifaWorldCup.3mf"),
    18: ("Caixa para cartas UNO Flip", "casa", True, "UNO_Fip(2).3mf"),
    20: ("Suporte cristal para comando PS5", "casa", False, "stand_for_psjy.3mf"),
    21: ("Suporte Senna", "decoracao", True, "senna-memory.3mf"),
    22: ("Fidget pequeno antistress", "brinquedos", False, None),
    23: ("Fidget grande antistress", "brinquedos", False, "spiral_cone_fidget_both_parts(2).3mf"),
    24: ("Porta-cromos preto pequeno Mundial 2026", "casa", True, "CAIXA_PORTA_FIGURINHA_COPA_2026.3mf"),
    26: ("Porta-chaves taça do Mundial", "porta-chaves", True, "WorldCup_keychain(2).3mf"),
    28: ("Porta-óculos panda", "casa", False, "Panda_Glasses_Tray.3mf"),
    29: ("Máscara abóbora Halloween", "decoracao", False, None),
    30: ("Porta-chaves laranja articulada", "porta-chaves", False, None),
    31: ("Porta-chaves cavaleiro mini", "porta-chaves", True, None),
    32: ("Porta-chaves coração", "porta-chaves", False, None),
    33: ("Máscara caveira Halloween", "decoracao", False, None),
    34: ("Porta-chaves polvo", "porta-chaves", False, None),
    35: ("Porta-chaves dragão articulado", "porta-chaves", False, None),
    36: ("Porta-chaves caveira com capuz", "porta-chaves", False, None),
    37: ("Porta-chaves esqueletos articulados", "porta-chaves", False, None),
    38: ("Porta-chaves Stanley Cup", "porta-chaves", True, None),
    39: ("Porta-chaves abóbora Halloween", "porta-chaves", False, None),
    40: ("Porta-chaves Pokébola", "porta-chaves", True, "pokeball keychain v2.3mf"),
    41: ("Porta-chaves bebé Darth Vader", "porta-chaves", True, None),
    42: ("Porta-chaves Creeper", "porta-chaves", True, None),
    43: ("Porta-chaves Grogu", "porta-chaves", True, None),
    44: ("Porta-chaves axolote flexível", "porta-chaves", False, None),
    45: ("Suporte para cápsulas", "casa", False, None),
    46: ("Porta-chaves dinossauro", "porta-chaves", False, None),
    47: ("Fidget estrela", "brinquedos", False, "STAR FIDGET W_HANDLE.3mf"),
    48: ("Porta-chaves caveira", "porta-chaves", False, None),
    49: ("Tabuada Pikachu", "brinquedos", True, "7-SCHEDA_TABELLINE_BAMBINI.3mf"),
    50: ("Caveira Halloween para vela", "decoracao", False, None),
    53: ("Presépio de Natal", "decoracao", False, None),
    54: ("Alien Xenomorph", "decoracao", True, None),
    55: ("Escultura Organic Bond", "decoracao", False, None),
    56: ("Meia-lua ghost", "decoracao", False, None),
    58: ("Cubo infinito", "brinquedos", False, None),
}
assert len(NOVOS) == 47, len(NOVOS)

def slug(t):
    t = t.lower()
    for a, b in zip("áàâãéêíóôõúçñ", "aaaaeeiooouc n"):
        t = t.replace(a, b)
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")[:40]

def foto(row, ficheiro):
    """Extrai a capa do designer do .3mf para public/images/produtos (JPEG)."""
    z = zipfile.ZipFile(os.path.join(STL, ficheiro))
    nomes = z.namelist()
    pick = next(n for n in ["Auxiliaries/.thumbnails/thumbnail_middle.png", "Auxiliaries/.thumbnails/thumbnail_3mf.png", "Metadata/plate_1.png"] if n in nomes)
    im = Image.open(io.BytesIO(z.read(pick))).convert("RGBA")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    bg.paste(im, (0, 0), im)
    bg.thumbnail((1200, 1200))
    rel = f"images/produtos/excel-l{row}.jpg"
    bg.save(os.path.join(SITE, "public", rel), "JPEG", quality=86, optimize=True, progressive=True)
    return rel

wb = load_workbook(XLS, data_only=True)
ws = wb["Calculadora 3D"]
linhas = {}
for r in range(4, 59):
    nome, gramas, k = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 11).value
    if nome and k is not None:
        linhas[r] = {"nome": str(nome).strip(), "g": gramas, "k": float(k)}

cat = json.load(open(base_catalog, encoding="utf-8"))
ids = {p["id"] for p in cat["products"]}
for c in [{"id": "porta-chaves", "name": "Porta-chaves"}, {"id": "brinquedos", "name": "Fidgets & brinquedos"}]:
    if c["id"] not in {x["id"] for x in cat["categories"]}:
        cat["categories"].append(c)

def opcao(row):
    l = linhas[row]
    peso = round(max(0.05, (l["g"] or 1000) / 1000), 3)
    return [{"id": "padrao", "label": "Padrão", "price": preco(l["k"]), "weight": peso}]

rel = {"atualizados": [], "novos_visiveis": [], "escondidos_marca": [], "escondidos_sem_foto": []}
for row, (pid, titulo) in EXISTENTES.items():
    p = next(x for x in cat["products"] if x["id"] == pid)
    p["title"] = titulo
    p["options"] = opcao(row)          # visibilidade e fotos ficam como o utilizador as deixou
    rel["atualizados"].append((row, pid, titulo, linhas[row]["k"], p["options"][0]["price"], p["visible"]))

for row, (titulo, categoria, marca, ficheiro) in NOVOS.items():
    pid = f"xl{row}-{slug(titulo)}"[:60]
    if pid in ids:
        continue  # já importado antes: não duplicar
    imagens = [foto(row, ficheiro)] if ficheiro else []
    visivel = not marca and bool(imagens)
    cat["products"].append({
        "id": pid, "title": titulo, "category": categoria, "visible": visivel, "images": imagens,
        "description": "", "options": opcao(row), "colors": [], "custom": False
    })
    item = (row, titulo, linhas[row]["k"], opcao(row)[0]["price"])
    rel["novos_visiveis" if visivel else ("escondidos_marca" if marca else "escondidos_sem_foto")].append(item)

json.dump(cat, open(out_catalog, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
json.dump(rel, open(out_report, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print({k: len(v) for k, v in rel.items()}, "produtos no catálogo:", len(cat["products"]))
